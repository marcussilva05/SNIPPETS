# python -m venv venv (cria venv se necessário)
# .\venv\Scripts\activate (ativa a venv)
# deactivate (desativa venv)
# pip install -r requirements.txt
# Start-Process "C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe" -ArgumentList "--remote-debugging-port=9222", "--user-data-dir=C:\edge-debug"
# cd C:\Users\vinic\Desktop\Python\WebScraping01
# python TESTES_SCRAPING.py

# sys
from colorama import init, Fore, Back, Style
from datetime import datetime
import datetime
import os
import platform
import pkgutil
import psutil
import requests
import shutil
import socket
import subprocess
import sys
import wmi

# leitura
from docx import Document
from openpyxl import load_workbook
import csv
import json
import re
import unicodedata
import tarfile 
import zipfile   

# debug
from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.edge.options import Options
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.common.action_chains import ActionChains
import pyautogui
import pygetwindow as gw
import uiautomation as auto
import time

init(autoreset=True) # reseta a cor no próximo print
# ============================================================
# VARIÁVEIS
# ============================================================
conectar = '1' # '1' conecta ou não conecta
excel = '0' # '1' lê excel ou rg teste = 179874871

# ============================================================
# FUNÇÃO popup()
# ============================================================
def popup(driver, timeout=10):
    print(Fore.YELLOW + "\n🔍 Verificando se há pop-up")
    
    # ===== VARIÁVEIS =====
    popup_encontrado = False
    popup_janela = None
    popup_titulo = None
    popup_body = None
    popup_close = None
    
    # ===== PESQUISA F12 =====
    try:
        esperar = WebDriverWait(driver, timeout)
        popup_janela = esperar.until(EC.presence_of_element_located((By.ID, "sedUiModalWrapper_1")))
        popup_encontrado = True
        print(Fore.GREEN + "✅ Pop-up encontrado!")
    except:
        print(Fore.YELLOW + "ℹ️ Nenhum pop-up de manifestação encontrado")
    
    # ===== SE ENCONTROU, PEGA OS ELEMENTOS =====
    if popup_encontrado:
        try:
            popup_titulo = driver.find_element(By.ID, "sedUiModalWrapper_1title")
            print(Fore.BLUE + f"   Título: {popup_titulo.text}")
        except:
            pass
        
        try:
            popup_body = driver.find_element(By.ID, "sedUiModalWrapper_1body")
            print(Fore.BLUE + f"   Texto: {popup_body.text[:100]}...")
        except:
            pass
        
        try:
            popup_close = driver.find_element(By.ID, "sedUiModalWrapper_1close")
            print(Fore.GREEN + "   ✅ Botão fechar encontrado!")
        except:
            pass
        
        # ===== FECHA O POP-UP =====
        if popup_close:
            try:
                popup_close.click()
                print(Fore.GREEN + "✅ Pop-up fechado!")
                time.sleep(1)
            except:
                print(Fore.YELLOW + "⚠️ Não foi possível fechar o pop-up")
    
    # ===== MOSTRA O RESULTADO =====
    print(Fore.CYAN + "\n📊 ELEMENTOS DO POP-UP:")
    print(Fore.WHITE + f"   Janela: {'✅' if popup_janela else '❌'}")
    print(Fore.WHITE + f"   Título: {'✅' if popup_titulo else '❌'}")
    print(Fore.WHITE + f"   Body: {'✅' if popup_body else '❌'}")
    print(Fore.WHITE + f"   Fechar: {'✅' if popup_close else '❌'}")
    
    return {
        'encontrado': popup_encontrado,
        'janela': popup_janela,
        'titulo': popup_titulo,
        'body': popup_body,
        'fechar': popup_close
    }
    
# ============================================================
# FUNÇÃO ir_para_cadastro_funcional
# ============================================================
def ir_para_cadastro_funcional(driver, timeout=10):
    """Navega até a página de Consulta Funcional"""
    
    print(Fore.YELLOW + "\n🔜 Navegando para Cadastro Funcional...")
    
    esperar = WebDriverWait(driver, timeout)
    
    try:
        # 1. Clica no hamburguer
        esperar.until(EC.element_to_be_clickable((By.ID, "decorHamburgerButton"))).click()
        print(Fore.GREEN + "✅ Hamburguer clicado!")
        time.sleep(0.5)
        
        # 2. Clica em Recursos Humanos
        esperar.until(EC.element_to_be_clickable(
            (By.XPATH, "//ul[@id='decorAsidePopup']//a[contains(text(), 'Recursos Humanos')]")
        )).click()
        print(Fore.GREEN + "✅ Recursos Humanos clicado!")
        time.sleep(0.5)
        
        # 3. Clica em Funcional
        esperar.until(EC.element_to_be_clickable(
            (By.XPATH, "//ul[@id='decorAsidePopup']//a[text()='Funcional ']") 
        )).click()
        print(Fore.GREEN + "✅ Funcional clicado!")
        time.sleep(0.5)
        
        # 4. Clica em Consulta Funcional
        esperar.until(EC.element_to_be_clickable(
            (By.XPATH, "//ul[@id='decorAsidePopup']//a[contains(text(), 'Consulta Funcional')]")
        )).click()
        print(Fore.GREEN + "✅ Consulta Funcional clicado!")
        time.sleep(2)
        
        print(Fore.BLUE + f"📍 URL: {driver.current_url}")
        return True
        
    except Exception as e:
        print(Fore.RED + f"❌ Erro: {e}")
        return False
    
# ============================================================
# FUNÇÃO ler_excel
# ============================================================
def ler_excel():
    # ------ LER EXCEL ------
    pasta_destino = os.chdir(r'C:\Users\vinic\Downloads')
    arquivo_excel = 'UDEMO_EDITADO.xlsx'
    numero_aba = 1  # Mude para 1, 2, 3, 4...
    coluna_indice = 1  # Mude para 0, 1, 2, 3...

    print(Fore.GREEN + Style.BRIGHT + f"\n🔎 Lendo arquivo '{arquivo_excel}' no endereço: '{os.getcwd()}'")

    if os.path.exists(arquivo_excel):
        print(Fore.GREEN + Style.BRIGHT + f"\n✅ Arquivo '{arquivo_excel}' existe!!")
        
    else:
        print(Fore.RED + Style.BRIGHT + f"\n❌ Arquivo '{arquivo_excel}' NÃO ENCONTRADO!")
        
    # ------ LER ABAS DO ARQUIVO EXCEL ------
        
    # Carrega o Excel
    excel_aberto = load_workbook(arquivo_excel, data_only=True, read_only=True)

    # Lista as abas
    abas = excel_aberto.sheetnames

    print(Fore.CYAN + f"\n📌 ABAS DISPONÍVEIS ({len(abas)}):")
    for i, nome in enumerate(abas, 1):
        print(Fore.WHITE + f"  {i}. {nome}")
        
    # ------ LER CABEÇALHO DA ABA SELECIONADA ------
    nome_aba = abas[numero_aba - 1]

    print(Fore.GREEN + f"\n✅ ABA SELECIONADA: {nome_aba}")

    # Carrega a aba
    aba = excel_aberto[nome_aba]
    linhas = list(aba.iter_rows(values_only=True))

    # Lê o cabeçalho
    if len(linhas) > 0:
        cabecalho = linhas[0]
        
        print(Fore.CYAN + f"\n📋 CABEÇALHO DA ABA '{nome_aba}':")
        print(Fore.WHITE + "-"*50)
        
        for i, coluna in enumerate(cabecalho):
            print(Fore.WHITE + f"  {i:2}. {coluna}")
        
        print(Fore.WHITE + "-"*50)
        print(Fore.YELLOW + f"Total de colunas: {len(cabecalho)}")
    else:
        print(Fore.RED + "❌ A planilha está vazia!")

    # ===== LER DADOS DA COLUNA PELO ÍNDICE =====
    print(Fore.CYAN + f"\n📄 DADOS DA COLUNA '{cabecalho[coluna_indice]}':")
    print(Fore.WHITE + "-"*50)

    dados = []
    for row in linhas[1:]:
        valor = row[coluna_indice] if row[coluna_indice] is not None else ''
        if valor:
            dados.append(valor)

    for i, valor in enumerate(dados[:10], 1):
        print(Fore.WHITE + f"  {i:3}. {valor}")

    # ===== AVISA SE TIVER MAIS DE 10 =====
    if len(dados) > 10:
        print(Fore.YELLOW + f"  ... e mais {len(dados) - 10} valores")

    print(Fore.WHITE + "-"*50)
    print(Fore.GREEN + f"✅ Total: {len(dados)} valores")
    
    excel_aberto.close()
    return dados
    
# ============================================================
# FUNÇÃO edge_debug()
# ============================================================
def edge_debug(fechar_edge_primeiro=False):   
     # ===== 1. FECHA EDGE SE SOLICITADO =====
    if fechar_edge_primeiro:
        print("🔄 Fechando Edge...")
        try:
            subprocess.run(["taskkill", "/f", "/im", "msedge.exe"], 
                          capture_output=True, text=True)
            print(Fore.GREEN + "✅ Edge fechado com sucesso!")
            time.sleep(2)
        except Exception as e:
            print(Fore.YELLOW + f"⚠️ Erro ao fechar Edge: {e}")
    else:
        print(Fore.BLUE + "ℹ️ Mantendo Edge aberto (não vai fechar)")
    
    # ===== 2. ABRE EDGE EM MODO DEBUG =====
    print("🔄 Abrindo Edge em modo debug...")
    edge_path = r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe"
    
    try:
        subprocess.Popen([
            edge_path, 
            "--remote-debugging-port=9222", 
            "--user-data-dir=C:\\edge-debug"
        ], shell=True)
        print(Fore.GREEN + "✅ Edge aberto em modo debug!")
        print(Fore.WHITE + "   📌 Porta: 9222")
        print(Fore.WHITE + "   📌 Perfil: C:\\edge-debug")
        time.sleep(3)
    except Exception as e:
        print(Fore.RED + f"❌ Erro ao abrir Edge: {e}")
        return None
    
    # ===== 3. CONECTA VIA SELENIUM =====
    print("🔄 Conectando via Selenium...")
    try:
        options = Options()
        options.add_experimental_option("debuggerAddress", "127.0.0.1:9222")
        
        driver = webdriver.Edge(options=options)
        print(Fore.GREEN + "✅ Conectado ao Edge com Selenium!")
        print(Fore.WHITE + f"   📌 URL atual: {driver.current_url}")
        print(Fore.WHITE + f"   📌 Título: {driver.title}")
        
        return driver
        
    except Exception as e:
        print(Fore.RED + f"❌ Erro ao conectar: {e}")
        print(Fore.YELLOW + "💡 Certifique-se de que o Edge está rodando com:")
        print(Fore.CYAN + '   --remote-debugging-port=9222')
        return None

# ============================================================
# CHAMAR edge_debug()
# ============================================================
if conectar == '1':
    driver = edge_debug(fechar_edge_primeiro=False)  # ← NÃO FECHA se True então ← FECHA E REABRE
    print(Fore.GREEN + Style.BRIGHT + f"✅ Conectado à URL: {driver.current_url}")
    print(Fore.GREEN + Style.BRIGHT + f"✅ Conectado ao site: {driver.title}")
    # ----- CONECTAR-SE A UM SITE ----- 
    url = 'https://sed.educacao.sp.gov.br/Inicio'
    esperar = WebDriverWait(driver, 30, poll_frequency=1) #verifica a cada 1s se está "clicável" até o máx de 30s    
    try: 
        driver.get(url)
        print(Fore.BLUE + Style.BRIGHT + f"\n🛜  site atual: {driver.current_url}")
        # ============================================================
        # CHAMAR popup()
        # ============================================================ 
        popup_info = popup(driver, timeout=5)
        # ============================================================
        # CHAMAR ir_para_cadastro_funcional()
        # ============================================================
        cadastro_info = ir_para_cadastro_funcional(driver, timeout=10)
    except:
        print(Fore.RED + Style.BRIGHT + "\n❌ Não conectou")
else:
    print(Fore.RED + Style.BRIGHT + f"✅ Não conectado no Edge em modo debug")
# ============================================================
# FUNÇÃO popup_servidor_nao_encontrado()
# ============================================================
def popup_servidor_nao_encontrado(driver, timeout=10):
    """
    Detecta o pop-up 'Servidor não encontrado!' procurando pelo texto.
    """
    try:
        print(Fore.YELLOW + "🔍 Procurando pop-up 'Servidor não encontrado'...")
        
        esperar = WebDriverWait(driver, timeout)
        
        # ===== PROCURA PELO TEXTO (XPATH) =====
        # Procura qualquer elemento que contenha "Servidor não encontrado"
        popup = esperar.until(
            EC.presence_of_element_located(
                (By.XPATH, "//*[contains(text(), 'Servidor não encontrado')]")
            )
        )
        
        if popup:
            print(Fore.GREEN + "✅ Pop-up 'Servidor não encontrado' detectado pelo texto!")
            
            # Pega o texto completo
            texto = popup.text
            print(Fore.YELLOW + f"📝 Texto: {texto}")
            
            # ===== TENTA FECHAR =====
            try:
                # Procura o botão Fechar
                btn = driver.find_element(By.XPATH, "//button[contains(text(), 'Fechar')]")
                btn.click()
                print(Fore.GREEN + "✅ Pop-up fechado!")
            except:
                try:
                    btn = driver.find_element(By.CLASS_NAME, "msg-button")
                    btn.click()
                    print(Fore.GREEN + "✅ Pop-up fechado!")
                except:
                    # Força via JavaScript
                    driver.execute_script("""
                        var popup = document.querySelector('.blockUI');
                        if (popup) popup.style.display = 'none';
                        var overlay = document.querySelector('.blockOverlay');
                        if (overlay) overlay.style.display = 'none';
                    """)
                    print(Fore.GREEN + "✅ Pop-up fechado via JavaScript!")
            
            return True
            
    except Exception as e:
        print(Fore.YELLOW + f"ℹ️ Nenhum pop-up 'Servidor não encontrado': {e}")
        return False
    
# ============================================================
# FUNÇÃO teste_rg (CORRIGIDA - COMPLETA)
# ============================================================
def teste_rg(driver, dados, timeout=15):
    print(Fore.YELLOW + "\n🪪 Testar RG")
    esperar = WebDriverWait(driver, timeout)
    
    try:
        # 1. Limpa o campo
        esperar.until(EC.element_to_be_clickable((By.ID, "btnLimpar"))).click()
        print(Fore.GREEN + "✅ Limpar clicado!")
        time.sleep(0.5)
        
        # 2. Clica no campo RG
        esperar.until(EC.element_to_be_clickable((By.ID, "txtNrRg"))).click()
        print(Fore.GREEN + "✅ RG clicado!")
        time.sleep(0.5)
        
        # 3. Digita o RG
        if dados:
            rg_para_digitar = str(dados[0]).strip()
            campo_digitar = driver.find_element(By.ID, "txtNrRg")
            campo_digitar.clear()
            campo_digitar.send_keys(rg_para_digitar)
            print(Fore.GREEN + f"\n✅ RG digitado: {rg_para_digitar}")
        else:
            print(Fore.RED + "❌ Lista de dados vazia!")
            return False
        
        # 4. Clica em Pesquisar
        esperar.until(EC.element_to_be_clickable((By.ID, "btnPesquisar"))).click()
        print(Fore.GREEN + "✅ Pesquisar clicado!")
        time.sleep(2)  # Aguarda a resposta
        
        # 5. VERIFICA POP-UP PELO TEXTO (NOVA FUNÇÃO)
        if popup_servidor_nao_encontrado(driver, timeout=10):
            print(Fore.RED + "❌ Servidor NÃO encontrado!")
            return False
        else:
            print(Fore.GREEN + "✅ Servidor ENCONTRADO!")
            return True
            
    except Exception as e:
        print(Fore.RED + f"❌ Erro: {e}")
        return False
    
# ============================================================
# 1. CHAMAR O EXCEL
# ============================================================
if excel == '1':
    dados = ler_excel()
    print(Fore.GREEN + f"\n✅ {len(dados)} RGs carregados do Excel!")
else:
    dados = ['179874871']
    print(Fore.YELLOW + f"\n⚠️ Usando RG de exemplo: {dados[0]}")

# ============================================================
# 2. CHAMAR teste_rg() (FORA DO IF/ELSE) RELACIONANDO A EXCEL
# ============================================================
if 'driver' in locals() and driver:
    escrever_rg = teste_rg(driver, dados, timeout=10)
    if escrever_rg:
        print(Fore.GREEN + "\n✅ RG testado com sucesso!")
    else:
        print(Fore.RED + "\n❌ Falha ao testar RG")
else:
    print(Fore.RED + "\n❌ Driver não disponível para testar RG")


    